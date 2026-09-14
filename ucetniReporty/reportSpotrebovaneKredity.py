#!/usr/bin/env python3

from re import S
import sys, os
import MySQLdb
import openpyxl
from dotenv import load_dotenv
from datetime import datetime, timedelta
import argparse

# Load environment variables from .env file
load_dotenv()

# MariaDB connection details from environment variables
USERDB_DB_HOST = os.environ.get('USERDB_DB_HOST', 'localhost')
USERDB_DB_PORT = int(os.environ.get('USERDB_DB_PORT', 3306))
USERDB_DB_USERNAME = os.environ.get('USERDB_DB_USERNAME')
USERDB_DB_PASSWORD = os.environ.get('USERDB_DB_PASSWORD')
USERDB_DB_NAME = os.environ.get('USERDB_DB_NAME')

if not all([USERDB_DB_USERNAME, USERDB_DB_PASSWORD, USERDB_DB_NAME]):
  print("MariaDB environment variables (USERDB_DB_USERNAME, USERDB_DB_PASSWORD, USERDB_DB_NAME) must be set.", file=sys.stderr)
  sys.exit(1)

try:
  udb_conn = MySQLdb.connect(
    user=USERDB_DB_USERNAME,
    password=USERDB_DB_PASSWORD,
    host=USERDB_DB_HOST,
    port=USERDB_DB_PORT,
    database=USERDB_DB_NAME,
    charset='utf8mb4'
  )
except MySQLdb.Error as e:
  print(f"Error connecting to MariaDB: {e}", file=sys.stderr)
  sys.exit(1)

udb_cursor = udb_conn.cursor()

wb = openpyxl.Workbook() # new Excel workbook

# Parse command line arguments for --year and --month
parser = argparse.ArgumentParser(description='Generuj měsíční účetní report podle roku a měsíce.')
parser.add_argument('--year', type=int, help='Rok pro report (např. 2026)')
parser.add_argument('--month', type=int, help='Měsíc pro report (1-12)')
parser.add_argument('--save-to-filename', type=str, help='Název XLSX souboru kam uložit výsledek')

args = parser.parse_args()

year = args.year
month = args.month

if args.year is None or args.month is None or args.save_to_filename is None:
    parser.print_help()
    exit()

#
# Přehled spotřebovaných kreditů - internet
#
query = """
SELECT uk.datum, -uk.castka as `částka Kč`,
u.id as UID,
CASE u.TypPravniFormyUzivatele_id WHEN 2 THEN u.firma_nazev ELSE CONCAT(coalesce(u.jmeno, ''), ' ', coalesce(u.prijmeni,'')) END as `jméno/název`,
concat(coalesce(u.ulice_cp,''), ', ', coalesce(u.mesto,'')) as adresa,
uk.poznamka as `poznámka`
FROM UzivatelskeKonto uk
LEFT JOIN Uzivatel u ON (uk.Uzivatel_id=u.id)
WHERE uk.druzstvo=1
AND year(uk.datum) = %u AND month(uk.datum) = %u
AND uk.castka != 0
AND (
	(uk.TypPohybuNaUctu_id = 4 AND uk.poznamka = '[Automaticka platba] hkfree.org internetove druzstvo')
	OR
	(uk.TypPohybuNaUctu_id = 8) -- Aktivace od [uid]
)
ORDER BY uk.datum, uk.datum_cas
"""

udb_cursor.execute(query % (year, month))
rows = udb_cursor.fetchall()

# Get column names
column_names = [desc[0] for desc in udb_cursor.description]

# Create worksheet
ws = wb.active
ws.title = "spotřebované kredity internet"

# Write headers
ws.append(column_names)

# Column widths
ws.column_dimensions['A'].width = 10 # datum
ws.column_dimensions['B'].width = 9  # castka_kc
ws.column_dimensions['C'].width = 9  # ucastnik_id
ws.column_dimensions['D'].width = 20 # ucastnik_nazev
ws.column_dimensions['E'].width = 30 # ucastnik_lokalita
ws.column_dimensions['F'].width = 60 # poznamka

for row_idx, row in enumerate(rows):
    # Write each row to the worksheet
    ws.append(row)

    # Formatting
    if row_idx >= 2:
        ws.cell(row=row_idx, column=1).number_format = 'YYYY-MM-DD'
        ws.cell(row=row_idx, column=2).number_format = '#,##0.00'



#
# Přehled spotřebovaných kreditů - televize
#
query = """
SELECT uk.datum, -uk.castka as `částka Kč`,
u.id as UID,
CASE u.TypPravniFormyUzivatele_id WHEN 2 THEN u.firma_nazev ELSE CONCAT(coalesce(u.jmeno, ''), ' ', coalesce(u.prijmeni,'')) END as `jméno/název`,
concat(coalesce(u.ulice_cp,''), ', ', coalesce(u.mesto,'')) as adresa,
uk.poznamka as `poznámka`
FROM UzivatelskeKonto uk
LEFT JOIN Uzivatel u ON (uk.Uzivatel_id=u.id)
WHERE uk.druzstvo=1
AND year(uk.datum) = %u AND month(uk.datum) = %u
AND uk.castka != 0
AND (
	(uk.TypPohybuNaUctu_id = 4 AND uk.poznamka = '[Automat] Aktivace služby Televize na 1 kalendářní měsíc')
	OR
	(uk.TypPohybuNaUctu_id = 4 AND uk.castka = -160 AND uk.poznamka like '%%televize%%') -- ruční odečty
	OR
	(uk.TypPohybuNaUctu_id = 7 AND uk.castka = 160 AND uk.poznamka like '%%televize%%') -- ruční vratky
)
ORDER BY uk.datum, uk.datum_cas
"""

udb_cursor.execute(query % (year, month))
rows = udb_cursor.fetchall()

# Get column names
column_names = [desc[0] for desc in udb_cursor.description]

# Create 2nd worksheet
ws = wb.create_sheet(title="spotřebované kredity televize")

# Write headers
ws.append(column_names)

# Column widths
ws.column_dimensions['A'].width = 10 # datum
ws.column_dimensions['B'].width = 9  # castka_kc
ws.column_dimensions['C'].width = 9  # ucastnik_id
ws.column_dimensions['D'].width = 20 # ucastnik_nazev
ws.column_dimensions['E'].width = 30 # ucastnik_lokalita
ws.column_dimensions['F'].width = 60 # poznamka

for row_idx, row in enumerate(rows):
    # Write each row to the worksheet
    ws.append(row)

    # Formatting
    if row_idx >= 2:
        ws.cell(row=row_idx, column=1).number_format = 'YYYY-MM-DD'
        ws.cell(row=row_idx, column=2).number_format = '#,##0.00'


# Save to XLSX file
output_file = args.save_to_filename
wb.save(output_file)

# Clean up
udb_cursor.close()
udb_conn.close()
